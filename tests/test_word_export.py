"""Offline Word title and section pagination regressions."""
import tempfile
import unittest
import json
from pathlib import Path

from docx import Document

from report_service.word_export import (
    CHAPTER_ORDER,
    _render_md_lines,
    add_heading_styled,
    build_compact_docx,
    md_to_docx,
    _infer_compact_sections,
)


class WordExportTests(unittest.TestCase):
    def test_native_numbered_report_keeps_all_body_in_navigable_sections(self):
        content='1. 区域基本盘：数值与来源保留。\n2. 三大产业方向的确定依据：原文。\n3. 区级政策工具一：完整条件。\n4. 待核实事项汇总：原文未知。'
        groups,order=_infer_compact_sections(content)
        self.assertEqual([x for key in order for x in groups[key]],[
            '1. 区域基本盘：数值与来源保留。','2. 三大产业方向的确定依据：原文。',
            '3. 区级政策工具一：完整条件。','4. 待核实事项汇总：原文未知。'])
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);(root/'09_compact_report.md').write_text(content)
            build_compact_docx('闵行区',root,root/'brief.docx')
            doc=Document(root/'brief.docx')
            body='\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('完整条件',body)
            self.assertTrue(any(p.style.name=='Heading 1' for p in doc.paragraphs))
            self.assertFalse(any(p.style.name=='List Number' for p in doc.paragraphs))
            self.assertIn('PAGE',doc.sections[0].footer._element.xml)
            self.assertIn('w:hyperlink',doc._element.xml)

    def test_main_report_includes_all_structured_entities_and_keeps_research_separate(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            for filename in CHAPTER_ORDER:
                (root/filename).write_text('ONLY-RESEARCH-ATTACHMENT '+filename)
            (root/'09_compact_report.md').write_text('1. 区域基本盘：COMPACT-ORIGINAL。')
            metadata={'stages':{'05a_target_enterprises_dir1.md':{'selected':[{'name':'独立企业甲','landing_status':'待核实','rationale':'完整适配条件','uncertainty':'投资意愿未知'}]},
                '07_matching_score.md':{'scores':[{'name':'独立企业甲','rank':1,'weighted_score':7.4,'dimensions':{'industry_fit':8,'risk':3}}]},
                '06b_fact_check.md':{'checks':[{'claim':'原始核验主张','verdict':'待核实','source':'https://gov.cn/one','cross_source':'https://gov.cn/two'}]}}}
            (root/'evidence.json').write_text(json.dumps(metadata))
            md_to_docx(root,root/'main.docx','闵行区')
            doc=Document(root/'main.docx');body='\n'.join(p.text for p in doc.paragraphs)
            for text in ('COMPACT-ORIGINAL','独立企业甲','完整适配条件','投资意愿未知','7.4','原始核验主张','https://gov.cn/two'):
                self.assertIn(text,body)
            self.assertNotIn('ONLY-RESEARCH-ATTACHMENT',body)
    def test_compact_cover_and_header_preserve_region_suffix(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / '09_compact_report.md').write_text('一、总体判断\n已核实内容。', encoding='utf-8')
            for city, expected in (
                ('苏州', '苏州市'),
                ('苏州市', '苏州市'),
                ('松江区', '松江区'),
                ('祁东县', '祁东县'),
                (' 苏州市 ', '苏州市'),
            ):
                with self.subTest(city=city):
                    output = root / 'compact.docx'
                    build_compact_docx(city, root, output)
                    document = Document(output)
                    self.assertEqual(document.paragraphs[1].text, expected + '精准招商作战报告')
                    self.assertEqual(document.sections[0].header.paragraphs[0].text,
                                     expected + '精准招商作战报告（精简版）')

    def test_full_cover_uses_same_region_rule(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for filename in CHAPTER_ORDER:
                (root / filename).write_text('## 已核实章节\n事实内容。', encoding='utf-8')
            for city, expected in (('宝鸡', '宝鸡市'), ('宝鸡市', '宝鸡市'), ('北仑区', '北仑区')):
                with self.subTest(city=city):
                    output = root / 'full.docx'
                    md_to_docx(root, output, city)
                    self.assertEqual(Document(output).paragraphs[1].text,
                                     expected + '精准招商作战报告')

    def test_plain_chinese_section_titles_keep_following_body(self):
        document = Document()
        titles = ['一、总体判断', '五、政策工具与时效', '十一、后续行动', '二．企业评分', '三. 产业链']
        lines = []
        for title in titles:
            lines.extend([title, '事实正文。'])
        _render_md_lines(document, lines)
        for index, title in enumerate(titles):
            heading, body = document.paragraphs[index * 2:index * 2 + 2]
            self.assertEqual(heading.text, title)
            self.assertTrue(heading.paragraph_format.keep_with_next)
            self.assertTrue(heading.runs[0].bold)
            self.assertIsNotNone(heading._p.find('w:pPr/w:keepNext', heading._p.nsmap))
            self.assertEqual(body.text, '事实正文。')
            self.assertIsNone(body.paragraph_format.keep_with_next)

    def test_existing_markdown_headings_keep_following_body(self):
        document = Document()
        _render_md_lines(document, ['## 二级标题', '正文', '### 三级标题', '正文', '#### 四级标题', '正文'])
        for paragraph in document.paragraphs[::2]:
            self.assertTrue(paragraph.paragraph_format.keep_with_next)
        direct = add_heading_styled(document, '章节标题', 1)
        self.assertTrue(direct.paragraph_format.keep_with_next)

    def test_numeric_lists_and_ordinary_chinese_remain_body(self):
        document = Document()
        _render_md_lines(document, ['1. 已核实事实', '五个产业方向需要比较。', '一、', '普通正文。'])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs],
                         ['已核实事实', '五个产业方向需要比较。', '一、', '普通正文。'])
        self.assertEqual(document.paragraphs[0].style.name, 'List Number')
        for paragraph in document.paragraphs:
            self.assertIsNone(paragraph.paragraph_format.keep_with_next)


if __name__ == '__main__':
    unittest.main()
