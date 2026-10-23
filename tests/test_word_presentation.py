"""Presentation fixes preserve source claims, section structure, and URL bytes."""
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from report_service.word_export import (
    _decision_body_content,
    _infer_compact_sections,
    _section_body_lines,
    _render_md_lines,
    build_compact_docx,
    md_to_docx,
)


class WordPresentationTests(unittest.TestCase):
    def test_matching_first_heading_is_rendered_once_with_navigation(self):
        content = '一、总体结论\n1. 未知不等于不存在。\n二、项目条件\n2. 仅为待核线索。'
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / '09_compact_report.md'
            source.write_text(content, encoding='utf-8')
            build_compact_docx('随州市', root, root / 'brief.docx')
            doc = Document(root / 'brief.docx')
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
            self.assertEqual(headings, ['目录', '总体结论', '项目条件'])
            self.assertEqual(source.read_text(encoding='utf-8'), content)
            body = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('1. 未知不等于不存在。', body)
            self.assertIn('2. 仅为待核线索。', body)
            self.assertIn('PAGEREF chapter_1', doc._element.xml)

    def test_dedup_does_not_remove_later_subheadings_or_similar_prose(self):
        cases = [
            (['## 条件', '正文', '### 条件', '后续正文'], '条件',
             ['正文', '### 条件', '后续正文']),
            (['## 条件补充', '正文'], '条件', ['## 条件补充', '正文']),
            (['条件：金额待核。', '## 条件'], '条件', ['条件：金额待核。', '## 条件']),
            (['1. 条件需要核查。'], '条件', ['1. 条件需要核查。']),
            (['', '<!-- SECTION: 条件 -->', '## 条件', '原文'], '条件',
             ['', '<!-- SECTION: 条件 -->', '原文']),
            (['二、条件', '原文'], '条件（续2）', ['原文']),
        ]
        for lines, label, expected in cases:
            with self.subTest(lines=lines):
                before = list(lines)
                self.assertEqual(_section_body_lines(lines, label), expected)
                self.assertEqual(lines, before)

    def test_main_title_changes_only_first_current_city_artifact_title(self):
        title = '随州市招商调研精简报告（截至2026年10月）'
        content = title + '\n一、总体结论\n1. 另见精简报告中的条件。\n' + title
        expected = content.replace(title, '随州市招商调研报告（截至2026年10月）', 1)
        self.assertEqual(_decision_body_content(content, '随州'), expected)
        self.assertEqual(_decision_body_content(content, '十堰市'), content)
        self.assertEqual(_decision_body_content('正文：' + content, '随州市'), '正文：' + content)
        self.assertEqual(_decision_body_content('引用《' + title + '》', '随州市'),
                         '引用《' + title + '》')

    def test_heading_hashes_are_content_without_a_separating_space(self):
        for title in ('C#', 'F#', 'https://host.test/page#', 'https://host.test/page#section'):
            with self.subTest(title=title):
                lines = ['## ' + title, '正文']
                sections, order = _infer_compact_sections('\n'.join(lines))
                self.assertEqual(order, [title])
                self.assertEqual(sections[title], lines)
                self.assertEqual(_section_body_lines(lines, title), ['正文'])
                self.assertEqual(_section_body_lines(lines, title.split('#', 1)[0]), lines)
        self.assertEqual(_section_body_lines(['## C# ##', '正文'], 'C#'), ['正文'])
        self.assertEqual(_infer_compact_sections('## C# ##\n正文')[1], ['C#'])

    def test_main_and_compact_use_distinct_titles_without_modifying_input(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            content = '随州市招商调研精简报告（截至2026年10月）\n一、总体结论\n1. 条件待核。'
            source = root / '09_compact_report.md'
            source.write_text(content, encoding='utf-8')
            evidence = {'stages': {'05a_target_enterprises_dir1.md': {
                'selected': [{'name': '测试企业', 'uncertainty': '投资意愿未知'}]}}}
            (root / 'evidence.json').write_text(json.dumps(evidence), encoding='utf-8')
            md_to_docx(root, root / 'main.docx', '随州市')
            build_compact_docx('随州市', root, root / 'brief.docx')
            main = Document(root / 'main.docx')
            brief = Document(root / 'brief.docx')
            self.assertIn('随州市招商调研报告（截至2026年10月）', [p.text for p in main.paragraphs])
            self.assertNotIn('随州市招商调研精简报告（截至2026年10月）', [p.text for p in main.paragraphs])
            self.assertIn('随州市招商调研精简报告（截至2026年10月）', [p.text for p in brief.paragraphs])
            for doc in (main, brief):
                self.assertEqual(sum(p.text == '总体结论' for p in doc.paragraphs), 1)
                self.assertIn('1. 条件待核。', [p.text for p in doc.paragraphs])
            self.assertEqual(source.read_text(encoding='utf-8'), content)
            self.assertEqual(json.loads((root / 'evidence.json').read_text()), evidence)

    def test_source_paragraphs_keep_together_without_changing_url_or_claim(self):
        url = 'https://www.mem.gov.cn/gk/gwgg/agwzlfl/tz_01/202606/t20260629_000001.shtml?x=1&y=2'
        lines = [f'待核事项：原文条件必须核实（{url}）。',
                 f'1. 编号条件（{url}）。', f'- 列表条件（{url}）。',
                 '| 来源 | 条件 |', '| --- | --- |', f'| {url} | 仍待核 |',
                 '普通正文保留自然分页。']
        doc = Document()
        _render_md_lines(doc, lines)
        for paragraph in doc.paragraphs:
            if url in paragraph.text:
                self.assertTrue(paragraph.paragraph_format.keep_together)
                self.assertIn(url, paragraph.text)
            elif paragraph.text == '普通正文保留自然分页。':
                self.assertIsNone(paragraph.paragraph_format.keep_together)
        cell = doc.tables[0].cell(1, 0).paragraphs[0]
        self.assertEqual(cell.text, url)
        self.assertTrue(cell.paragraph_format.keep_together)


if __name__ == '__main__':
    unittest.main()
