"""Word navigation must never infer headings from words inside assertions."""
import copy
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from report_service.word_export import (
    _display_direction_labels,
    _infer_compact_sections,
    _split_compact_sections,
    build_compact_docx,
    md_to_docx,
)


class WordNavigationTests(unittest.TestCase):
    def test_summary_keywords_do_not_become_a_heading(self):
        lines = [
            '1. 本报告覆盖经济、产业链缺口、政策工具与行动，所有数据待核实。',
            '2. 经济底盘：原始数值与条件。',
            '3. 某文件提及产业链缺口：这是原文引用，不是新章节。',
            '4. 产业链环节分布方面，原始本地未知信息。',
            '5. 断点与卡点方面，原始断点证据。',
            '6. 国家产业政策方面，保留适用对象。',
        ]
        sections, order = _infer_compact_sections('\n'.join(lines))
        self.assertEqual(order, ['概览', '区域概况', '产业链与断点', '政策工具与条件'])
        self.assertEqual([line for key in order for line in sections[key]], lines)
        self.assertEqual(sections['区域概况'][-1], lines[2])

    def test_explicit_titles_and_revisited_sections_keep_source_order(self):
        lines = ['## 政策工具', '全文甲', '一、企业说明', '全文乙', '## 政策工具', '全文丙']
        sections, order = _infer_compact_sections('\n'.join(lines))
        self.assertEqual(order, ['政策工具', '企业说明', '政策工具（续2）'])
        self.assertEqual([line for key in order for line in sections[key]], lines)

    def test_empty_input_has_no_phantom_contents_entry(self):
        self.assertEqual(_infer_compact_sections(''), ({}, []))

    def test_tagged_duplicate_sections_preserve_every_line_and_preamble(self):
        lines = ['原文引言', '<!-- SECTION: 评分 -->', '第一个评分。',
                 '<!-- SECTION: 评分 -->', '第二个评分。', '']
        sections, order = _split_compact_sections('\n'.join(lines))
        self.assertEqual(order, ['评分', '评分（续2）'])
        self.assertEqual(sections['_preamble'] + [line for key in order for line in sections[key]], lines)

    def test_direction_mapping_is_exact_and_never_rewrites_urls_or_longer_ids(self):
        mapping = {'dir1': '能源装备', 'dir2': '航空航天', 'dir3': '生物医药'}
        text = '闵行dir1材料；dir2与dir3。https://host/dir1/a dir10 source_dir1.md xdir1'
        self.assertEqual(_display_direction_labels(text, mapping),
                         '闵行能源装备材料；航空航天与生物医药。https://host/dir1/a dir10 source_dir1.md xdir1')
        self.assertEqual(_display_direction_labels('dir3待核。', {}), 'dir3待核。')

    def test_direction_id_before_its_saved_label_does_not_duplicate_that_label(self):
        mapping = {'dir1': '能源装备', 'dir2': '航空航天', 'dir3': '生物医药'}
        self.assertEqual(_display_direction_labels('与dir2航空航天锻件环节、dir3生物医药相关；dir1待核。', mapping),
                         '与航空航天锻件环节、生物医药相关；能源装备待核。')
        self.assertEqual(_display_direction_labels('dir2航空航天航空航天、dir2航空方向。', mapping),
                         '航空航天航空航天、航空航天航空方向。')

    def test_main_contents_nested_bookmarks_and_direction_labels_do_not_mutate_evidence(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            text = '1. 经济底盘：原文数值。\n2. 产业链环节分布方面，未知。\n3. 国家产业政策方面，条件待核。'
            (root / '09_compact_report.md').write_text(text)
            evidence = {'stages': {
                '02_industry_direction.md': {'directions': [{'id': 'dir1', 'name': '能源装备'}]},
                '05a_target_enterprises_dir1.md': {'selected': [{'name': '甲公司',
                    'rationale': 'dir1适配依据，投资意愿未知。',
                    'url': 'https://host/dir1/source', 'uncertainty': 'dir1供货待核。'}]},
                '07_matching_score.md': {'scores': []}, '06b_fact_check.md': {'checks': []}}}
            before = copy.deepcopy(evidence)
            (root / 'evidence.json').write_text(json.dumps(evidence))
            md_to_docx(root, root / 'main.docx', '闵行区')
            doc = Document(root / 'main.docx')
            body = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('1. 经济底盘：原文数值。', body)
            self.assertIn('能源装备适配依据，投资意愿未知。', body)
            self.assertIn('https://host/dir1/source', body)
            self.assertNotIn('dir1供货', body)
            self.assertEqual(json.loads((root / 'evidence.json').read_text()), before)
            anchors = {item.get(qn('w:name')) for item in doc._element.xpath('.//w:bookmarkStart')}
            links = {item.get(qn('w:anchor')) for item in doc._element.xpath('.//w:hyperlink')}
            self.assertTrue(links <= anchors)
            self.assertTrue({'contents', 'chapter_0', 'body_0', 'direction_1'} <= anchors)
            self.assertIn('PAGEREF body_0', doc._element.xml)
            self.assertIn('w:updateFields', doc.settings.element.xml)

    def test_compact_keeps_original_numbering_and_conditions_in_source_order(self):
        text = '1. 本报告产业链缺口仅为待核线索。\n2. 经济底盘：原始数据。\n3. 国家产业政策方面，资格必须先核。\n4. 综合结论：未知不等于不存在。'
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / '09_compact_report.md').write_text(text)
            build_compact_docx('闵行区', root, root / 'brief.docx')
            doc = Document(root / 'brief.docx')
            body = '\n'.join(p.text for p in doc.paragraphs)
            positions = [body.index(line) for line in text.splitlines()]
            self.assertEqual(positions, sorted(positions))
            self.assertFalse(any(p.style.name == 'List Number' for p in doc.paragraphs))


if __name__ == '__main__':
    unittest.main()
